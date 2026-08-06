"""Add an AFD-vs-colocated-with-TBO ratio column to a finished workbook.

The colocated arm as billed puts every all-to-all on the critical path. This
adds the opposite extreme: a colocated arm where two-batch overlap hides all of
it, ``step = max(compute, comm)``. The AFD arm is left exactly as it is -- its
pipeline already overlaps A against F, so both arms are then credited with the
overlap their design allows.

The winning colocated configuration is re-selected under the overlap, not
merely rescaled. Hiding the collectives rewards wide-EP layouts that the serial
billing penalised, so the config that wins with TBO is often not the one that
wins without it; rescaling the old winner would understate the colocated arm
and flatter AFD.

Usage: add_tbo_column.py <workbook.xlsx> <tbo.csv>
"""
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

WB, TBO = sys.argv[1], sys.argv[2]

GREEN = PatternFill("solid", fgColor="C6EFCE")
HEAD = Font(bold=True, size=9)
BODY = Font(size=9)

NOTE = ("The same AFD point divided by a colocated arm whose step time is "
        "max(compute, communication) instead of compute + communication -- the ceiling of "
        "two-batch overlap, where every collective is hidden at no cost. The colocated "
        "configuration and batch are re-optimised under that ceiling, because hiding the "
        "collectives rewards layouts that serial billing penalised. The AFD arm is "
        "unchanged; its A/F pipeline already overlaps. The ceiling is not attainable: it "
        "ignores that halving the batch shrinks every GEMM, ignores the second set of "
        "activation buffers, and ignores the sync points a real implementation keeps. "
        "colo_tbo_speedup is how much the overlap is worth at that point.")

tbo = pd.read_csv(TBO, low_memory=False)
tbo["mtp_nextn"] = tbo.mtp_nextn.astype(str)
tbo["attn"] = tbo["attn"].replace({"MSA (linear)": "MSA"})
tbo["step_now"] = tbo.compute_ms + tbo.comm_ms
tbo["factor"] = tbo.step_now / tbo[["compute_ms", "comm_ms"]].max(axis=1)
tbo["tps_tbo"] = tbo.tokens_s_gpu * tbo.factor
tbo["tpot_tbo"] = tbo.tpot_ms / tbo.factor
tbo["user_tbo"] = tbo.tokens_s_user * tbo.factor

# ``attn`` is not part of the key: the workbook labels it with the head ratio
# ("GQA 64:4") while the sweep writes the bare family, and the model already
# determines it.
CELL = ["model", "precision", "isl", "osl", "mtp_nextn"]
GROUPS = {k: g for k, g in tbo.groupby(CELL)}

# Budget sheet: same SLA the budget sheet itself applies.
BUDGET_SLA = 50.0


def best(sub, sla=None, speed=None, max_unit=None):
    ok = sub
    if sla is not None:
        ok = ok[ok.tpot_tbo <= sla]
    if speed is not None:
        ok = ok[ok.user_tbo >= speed]
    if max_unit is not None:
        ok = ok[ok.unit_gpus <= max_unit]
    return ok.loc[ok.tps_tbo.idxmax()] if len(ok) else None


wb = load_workbook(WB)
added = missing = 0

for ws in wb.worksheets:
    if ws.title == "Configuration":
        continue
    h = [c.value for c in ws[2]]
    if "afd_tpsgpu" not in h:
        continue
    ix = {n: h.index(n) for n in h if n}
    budget_sheet = "budget_gpus" in ix
    prec_col = ix.get("precision", ix.get("moe_precision"))
    sla_col = ix.get("sla_tpot_ms")
    spd_col = ix.get("target_tokps_per_user")

    base = ws.max_column
    hdrs = ["colo_tbo_cfg", "colo_tbo_batch", "colo_tbo_speedup"]
    hdrs += (["colo_tbo_replicas", "colo_tbo_total_tokens_s", "ratio_vs_colo_tbo"]
             if budget_sheet else ["colo_tbo_tpsgpu", "ratio_vs_colo_tbo"])
    for j, name in enumerate(hdrs):
        c = ws.cell(row=2, column=base + 1 + j, value=name)
        c.font = HEAD
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = 16

    for r in range(3, ws.max_row + 1):
        get = lambda c: ws.cell(row=r, column=c + 1).value
        if get(ix["model"]) is None:
            continue
        key = (get(ix["model"]), get(prec_col), get(ix["isl"]),
               get(ix["osl"]), str(get(ix["mtp_nextn"])))
        sub = GROUPS.get(key)
        if sub is None:
            missing += 1
            continue

        if budget_sheet:
            b = get(ix["budget_gpus"])
            afd_total = get(ix["afd_total_tokens_s"])
            w = best(sub, sla=BUDGET_SLA, max_unit=b)
            if w is None or afd_total in (None, ""):
                missing += 1
                continue
            reps = int(b // int(w.unit_gpus))
            total = float(w.tps_tbo) * int(w.unit_gpus) * reps
            vals = [w.config, int(w.batch_per_unit), round(float(w.factor), 3),
                    reps, round(total, 1), round(float(afd_total) / total, 4)]
        else:
            afd = get(ix["afd_tpsgpu"])
            sla = get(sla_col) if sla_col is not None else None
            spd = get(spd_col) if spd_col is not None else None
            w = best(sub, sla=sla, speed=spd)
            if w is None or afd in (None, ""):
                missing += 1
                continue
            vals = [w.config, int(w.batch_per_unit), round(float(w.factor), 3),
                    round(float(w.tps_tbo), 2), round(float(afd) / float(w.tps_tbo), 4)]

        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=base + 1 + j, value=v)
            c.font = BODY
            c.alignment = Alignment(horizontal="center")
        if vals[-1] > 1:
            ws.cell(row=r, column=base + len(vals)).fill = GREEN
        added += 1

for ws in wb.worksheets:
    if ws.title != "Configuration":
        continue
    row = ws.max_row + 2
    ws.cell(row=row, column=1, value="ratio_vs_colo_tbo").font = HEAD
    c = ws.cell(row=row, column=2, value=NOTE)
    c.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(WB)
print(f"{WB.split('/')[-1]}: 写入 {added} 行, 无匹配 {missing}")
