"""Add a fixed-GPU-budget sheet to an existing workbook, touching nothing else.

The result sheets answer "per GPU, which arrangement is better", which already
assumes each side is replicated to fill whatever budget exists. That hides two
things a purchase decision needs: the absolute throughput at a given budget,
and whether a deployment unit divides that budget without stranding GPUs. This
sheet fixes the budget, lets each side pick its own unit and replicate, and
reports the totals plus any GPUs left over.
"""
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WB, CSV = sys.argv[1], sys.argv[2]
SHEET = "GPU_Budget"
BUDGETS = [8, 16, 32, 64, 72]
SLA = 50
OSL = 1024

GREEN = PatternFill("solid", fgColor="C6EFCE")
HDR = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
NOTE_FONT = Font(italic=True, size=9, color="595959")

d = pd.read_csv(CSV, low_memory=False)
d = d[(d.tokens_s_gpu.notna()) & (d.osl == OSL)]
d["mtp_nextn"] = d.mtp_nextn.astype(str)
d["attn"] = d["attn"].replace({"MSA (linear)": "MSA"})
KEY = ["model", "attn", "precision", "isl", "osl", "mtp_nextn"]


def pick(g, budget):
    """Best replicated deployment for one arm inside the budget."""
    g = g[(g.tpot_ms <= SLA) & (g.unit_gpus <= budget)]
    if not len(g):
        return None
    reps = (budget // g.unit_gpus).astype(int)
    g = g.assign(reps=reps)
    g = g[g.reps >= 1]
    if not len(g):
        return None
    g = g.assign(total=g.tokens_s_gpu * g.unit_gpus * g.reps)
    return g.loc[g.total.idxmax()]


rows = []
for key, g in d.groupby(KEY):
    for b in BUDGETS:
        a, c = pick(g[g.arm == "afd"], b), pick(g[g.arm == "colocated"], b)
        if a is None and c is None:
            continue
        r = dict(zip(KEY, key))
        r["budget_gpus"] = b
        r["moe_kernel"] = g.moe_kernel.iloc[0]
        r["sla_tpot_ms"] = SLA
        if a is not None:
            used = int(a.unit_gpus * a.reps)
            r.update(afd_split=a.config, afd_unit_gpus=int(a.unit_gpus),
                     afd_replicas=int(a.reps), afd_gpus_used=used,
                     afd_gpus_stranded=b - used, afd_batch=int(a.batch_per_unit),
                     afd_tpot_ms=round(a.tpot_ms, 2),
                     afd_tpsgpu=round(a.tokens_s_gpu, 1),
                     afd_total_tokens_s=round(float(a.total), 1))
        if c is not None:
            used = int(c.unit_gpus * c.reps)
            r.update(colo_cfg=c.config, colo_unit_gpus=int(c.unit_gpus),
                     colo_replicas=int(c.reps), colo_gpus_used=used,
                     colo_gpus_stranded=b - used, colo_batch=int(c.batch_per_unit),
                     colo_tpot_ms=round(c.tpot_ms, 2),
                     colo_tpsgpu=round(c.tokens_s_gpu, 1),
                     colo_total_tokens_s=round(float(c.total), 1))
        if a is not None and c is not None and float(c.total) > 0:
            r["ratio"] = round(float(a.total) / float(c.total), 3)
        rows.append(r)

df = pd.DataFrame(rows)
lead = ["model", "attn", "precision", "isl", "osl", "mtp_nextn", "budget_gpus",
        "sla_tpot_ms", "moe_kernel"]
df = df[[c for c in lead if c in df.columns] + [c for c in df.columns if c not in lead]]
df = df.rename(columns={"precision": "moe_precision"})
df = df.sort_values(["model", "moe_precision", "isl", "mtp_nextn", "budget_gpus"])

wb = load_workbook(WB)
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws = wb.create_sheet(SHEET)
ws["A1"] = (f"Fixed GPU budget, generation length 1024, {SLA} ms time-per-output-token limit. "
            "Each side picks its own deployment unit and replicates as many times as the budget "
            "allows; the totals are summed over replicas and gpus_stranded is what the budget "
            "could not fill. Rows where the disaggregated total exceeds the colocated total are "
            "shaded.")
ws["A1"].font = NOTE_FONT
cols = list(df.columns)
for j, c in enumerate(cols, 1):
    cell = ws.cell(row=2, column=j, value=c)
    cell.fill, cell.font = HDR, HDR_FONT
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
for i, (_, r) in enumerate(df.iterrows(), 3):
    green = pd.notna(r.get("ratio")) and r["ratio"] > 1.0
    for j, c in enumerate(cols, 1):
        v = r[c]
        cell = ws.cell(row=i, column=j, value=None if pd.isna(v) else v)
        if green:
            cell.fill = GREEN
for j, c in enumerate(cols, 1):
    w = max([len(str(c))] + [len(str(v)) for v in df[c].head(300)])
    ws.column_dimensions[get_column_letter(j)].width = min(max(w + 2, 9), 24)
ws.freeze_panes = "H3"
wb.save(WB)
n_green = int((df.ratio > 1).sum()) if "ratio" in df else 0
print(f"{WB.split('/')[-1]}: 新增 {SHEET}, {len(df)} 行 x {len(cols)} 列, ratio>1 {n_green} 行")
