"""Recompute batch_limited_by in place, touching nothing else.

The original label compared the winning batch only against the largest batch
that fits HBM, so a batch the sheet's own latency or speed constraint had ruled
out was reported as ``compute`` -- reading as saturated compute when the limit
was really the constraint. This rewrites those cells (and the Configuration
line that defines them) and leaves every other cell, style and sheet untouched.
"""
import sys

import pandas as pd
from openpyxl import load_workbook

WB, CSV = sys.argv[1], sys.argv[2]

NEW_DEF = ("What stopped the winning batch from growing. capacity: no larger batch fits in "
           "HBM. latency: larger batches fit but breach the limit in force on that sheet. "
           "compute: larger batches both fit and meet the limit, yet do not raise throughput. "
           "grid: the batch grid ran out of rungs.")

d = pd.read_csv(CSV, low_memory=False)
d = d[d.tokens_s_gpu.notna()]
d["mtp_nextn"] = d.mtp_nextn.astype(str)
d["attn"] = d["attn"].replace({"MSA (linear)": "MSA"})
BATCH_MAX = int(d.batch_per_unit.max())

# (model, attn, precision, isl, osl, mtp, arm, config) -> feasible batch rows
GROUPS = {}
for key, g in d.groupby(["model", "attn", "precision", "isl", "osl", "mtp_nextn", "arm", "config"]):
    GROUPS[key] = g[["batch_per_unit", "tpot_ms", "tokens_s_user"]]


def verdict(key, batch, sla, speed):
    g = GROUPS.get(key)
    if g is None or not len(g):
        return None
    top_feasible = int(g.batch_per_unit.max())
    ok = g
    if sla is not None:
        ok = ok[ok.tpot_ms <= sla]
    if speed is not None:
        ok = ok[ok.tokens_s_user >= speed]
    top_allowed = int(ok.batch_per_unit.max()) if len(ok) else top_feasible
    if int(batch) < top_allowed:
        return "compute"
    if top_allowed < top_feasible:
        return "latency"
    return "capacity" if top_feasible < BATCH_MAX else "grid"


wb = load_workbook(WB)
changed = 0
unmatched = 0
before, after = {}, {}

for ws in wb.worksheets:
    if ws.title == "Configuration":
        for row in ws.iter_rows():
            if row[0].value == "batch_limited_by":
                row[1].value = NEW_DEF
                changed += 1
        continue

    h = [c.value for c in ws[2]]
    if "afd_batch_limited_by" not in h:
        continue
    ix = {n: h.index(n) for n in h if n}
    sla_col = ix.get("sla_tpot_ms")
    spd_col = ix.get("target_tokps_per_user")
    prec_col = ix.get("moe_precision", ix.get("precision"))

    for r in range(3, ws.max_row + 1):
        get = lambda c: ws.cell(row=r, column=c + 1).value
        if get(ix["model"]) is None:
            continue
        sla = get(sla_col) if sla_col is not None else None
        spd = get(spd_col) if spd_col is not None else None
        ident = (get(ix["model"]), get(ix["attn"]), get(prec_col), get(ix["isl"]),
                 get(ix["osl"]), str(get(ix["mtp_nextn"])))
        for arm, cfg_col, bs_col, lim_col in (
                ("afd", "afd_split", "afd_batch", "afd_batch_limited_by"),
                ("colocated", "colo_cfg", "colo_batch", "colo_batch_limited_by")):
            if lim_col not in ix or cfg_col not in ix or bs_col not in ix:
                continue
            cfg, bs = get(ix[cfg_col]), get(ix[bs_col])
            cell = ws.cell(row=r, column=ix[lim_col] + 1)
            if cfg is None or bs is None or cell.value in (None, ""):
                continue
            v = verdict(ident + (arm, cfg), bs, sla, spd)
            if v is None:
                unmatched += 1
                continue
            before[cell.value] = before.get(cell.value, 0) + 1
            after[v] = after.get(v, 0) + 1
            if v != cell.value:
                cell.value = v          # value only; fill and font stay as they are
                changed += 1

wb.save(WB)
print(f"{WB.split('/')[-1]}")
print(f"   改写 {changed} 个单元格, 未匹配 {unmatched}")
print(f"   改前: {dict(sorted(before.items()))}")
print(f"   改后: {dict(sorted(after.items()))}")
